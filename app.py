#!/usr/bin/env python3
"""Offline classroom attendance server (Flask + SQLite, pure-Python, Termux friendly).

v2 highlights
- Face matching happens on the server; descriptors are never sent back to the browser.
- Admin PIN (PBKDF2), idle timeout, CSRF token, Host/Origin checks, security headers.
- Multi-sample enrolment, duplicate-face detection, ambiguity margin, multi-frame confirmation.
- Late rule, excused absences, holidays/school calendar, real monthly attendance rate.
- Audit log, backup, retention purge, CSV-injection-safe exports.
- Automatic, non-destructive migration of the v1 database.

v2.1 additions
- Calibration tab: live face-size / distance / match readouts while enrolling (admin-only match probe).
- Long-range detection (up to ~1 m): larger detector input, 720p camera, near+far enrolment samples.
- Excel (.xlsx) downloads of the daily report and the monthly summary (summary, day-by-day, per-day totals).
"""
import csv
import hashlib
import hmac
import io
import json
import logging
import math
import os
import secrets
import sqlite3
import sys
import tempfile
import threading
import time
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from functools import wraps
from pathlib import Path

from flask import Flask, Response, jsonify, render_template, request, session
from werkzeug.exceptions import HTTPException

BASE = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get("ATTENDANCE_DATA_DIR", BASE))
DB_PATH = DATA_DIR / "attendance.db"
KEY_PATH = DATA_DIR / ".secret_key"
HOST, PORT = "127.0.0.1", 8765
ALLOWED_HOSTS = {f"127.0.0.1:{PORT}", f"localhost:{PORT}"}
ALLOWED_ORIGINS = {f"http://{h}" for h in ALLOWED_HOSTS}
DESCRIPTOR_LEN = 128
SCHEMA_VERSION = 2
ADMIN_IDLE_SECONDS = 10 * 60
MIN_PIN_LEN = 6

log = logging.getLogger("attendance")


# --------------------------------------------------------------------------- app
def _load_secret_key():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if KEY_PATH.exists():
        return KEY_PATH.read_bytes()
    key = secrets.token_bytes(32)
    fd = os.open(KEY_PATH, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(fd, "wb") as f:
        f.write(key)
    return key


app = Flask(__name__)
app.secret_key = _load_secret_key()
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Strict",
    SESSION_COOKIE_NAME="attendance_session",
    PERMANENT_SESSION_LIFETIME=timedelta(hours=12),
    MAX_CONTENT_LENGTH=1024 * 1024,
    JSON_SORT_KEYS=False,
)


def now():
    """Single clock source so tests can freeze time."""
    return datetime.now()


# ---------------------------------------------------------------------- database
@contextmanager
def db():
    con = sqlite3.connect(DB_PATH, timeout=10)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys=ON")
    try:
        yield con
        con.commit()
    except Exception:
        con.rollback()
        raise
    finally:
        con.close()


def _columns(con, table):
    return {r["name"] for r in con.execute(f"PRAGMA table_info({table})")}


def parse_samples(raw):
    """v1 stored one flat descriptor; v2 stores a list of descriptors."""
    data = json.loads(raw)
    if data and isinstance(data[0], (int, float)):
        return [data]
    return data


def init_db():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    with db() as con:
        try:
            con.execute("PRAGMA journal_mode=WAL")
        except sqlite3.DatabaseError:
            pass
        con.execute("""CREATE TABLE IF NOT EXISTS students (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            student_no TEXT UNIQUE NOT NULL,
            name TEXT NOT NULL,
            descriptor TEXT NOT NULL,
            active INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL)""")
        con.execute("""CREATE TABLE IF NOT EXISTS attendance (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            student_id INTEGER NOT NULL,
            attend_date TEXT NOT NULL,
            time_in TEXT NOT NULL,
            method TEXT NOT NULL DEFAULT 'face',
            UNIQUE(student_id, attend_date),
            FOREIGN KEY(student_id) REFERENCES students(id))""")
        # additive migrations (safe to run repeatedly, never drops data)
        if "section" not in _columns(con, "students"):
            con.execute("ALTER TABLE students ADD COLUMN section TEXT NOT NULL DEFAULT ''")
        if "updated_at" not in _columns(con, "students"):
            con.execute("ALTER TABLE students ADD COLUMN updated_at TEXT NOT NULL DEFAULT ''")
        if "status" not in _columns(con, "attendance"):
            con.execute("ALTER TABLE attendance ADD COLUMN status TEXT NOT NULL DEFAULT 'present'")
        if "note" not in _columns(con, "attendance"):
            con.execute("ALTER TABLE attendance ADD COLUMN note TEXT NOT NULL DEFAULT ''")
        con.execute("CREATE INDEX IF NOT EXISTS idx_att_date ON attendance(attend_date)")
        con.execute("""CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT NOT NULL)""")
        con.execute("""CREATE TABLE IF NOT EXISTS holidays (
            day TEXT PRIMARY KEY, reason TEXT NOT NULL DEFAULT '')""")
        con.execute("""CREATE TABLE IF NOT EXISTS audit_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT NOT NULL,
            action TEXT NOT NULL, detail TEXT NOT NULL DEFAULT '')""")
        # v1 -> v2 descriptor format: flat list -> list of samples
        for r in con.execute("SELECT id, descriptor FROM students").fetchall():
            try:
                samples = parse_samples(r["descriptor"])
            except (ValueError, TypeError):
                continue
            if r["descriptor"] != json.dumps(samples):
                con.execute("UPDATE students SET descriptor=? WHERE id=?", (json.dumps(samples), r["id"]))
        con.execute(f"PRAGMA user_version={SCHEMA_VERSION}")
    try:
        os.chmod(DB_PATH, 0o600)
    except OSError:
        pass


def audit(con, action, detail=""):
    con.execute("INSERT INTO audit_log(ts, action, detail) VALUES(?,?,?)",
                (now().isoformat(timespec="seconds"), action, detail[:500]))


# ---------------------------------------------------------------------- settings
SETTINGS = {
    # key: (kind, default, lo, hi)
    "class_start": ("time", "08:00"),
    "late_after_min": ("int", 15, 0, 240),
    "threshold": ("float", 0.48, 0.30, 0.60),
    "margin": ("float", 0.06, 0.0, 0.20),
    "confirm_frames": ("int", 3, 1, 8),
    "require_liveness": ("bool", False),
    "school_days": ("days", "1,2,3,4,5"),
    # long-range detection (read by the browser): detector input size and smallest accepted face width
    "detect_size": ("int", 512, 224, 608),
    "min_face_pct": ("float", 8.0, 3.0, 30.0),
}


def coerce_setting(key, value):
    kind, default, *rng = SETTINGS[key]
    if kind == "time":
        datetime.strptime(str(value), "%H:%M")
        return str(value)
    if kind == "int":
        v = int(value)
        if key == "detect_size":
            v = int(round(v / 32.0)) * 32  # TinyFaceDetector needs a multiple of 32
    elif kind == "float":
        v = float(value)
    elif kind == "bool":
        return "1" if value in (True, 1, "1", "true", "True") else "0"
    elif kind == "days":
        days = sorted({int(x) for x in (value if isinstance(value, list) else str(value).split(",")) if str(x).strip()})
        if not days or any(d < 1 or d > 7 for d in days):
            raise ValueError("school_days must be ISO weekdays 1-7 (Mon=1)")
        return ",".join(map(str, days))
    else:
        raise ValueError("unknown setting kind")
    if not math.isfinite(v) or v < rng[0] or v > rng[1]:
        raise ValueError(f"{key} must be between {rng[0]} and {rng[1]}")
    return str(v)


def _typed(key, raw):
    kind = SETTINGS[key][0]
    if kind == "int":
        return int(raw)
    if kind == "float":
        return float(raw)
    if kind == "bool":
        return raw == "1" or raw is True
    if kind == "days":
        return {int(x) for x in str(raw).split(",") if x}
    return raw


def get_settings(con):
    vals = {k: v[1] for k, v in SETTINGS.items()}
    for r in con.execute("SELECT key, value FROM settings"):
        if r["key"] in SETTINGS:
            vals[r["key"]] = r["value"]
    return {k: _typed(k, v) for k, v in vals.items()}


def settings_for_api(s):
    out = dict(s)
    out["school_days"] = sorted(s["school_days"])
    return out


# -------------------------------------------------------------------------- auth
def hash_pin(pin, salt=None, iters=200_000):
    salt = salt or secrets.token_bytes(16)
    dk = hashlib.pbkdf2_hmac("sha256", pin.encode(), salt, iters)
    return f"pbkdf2${iters}${salt.hex()}${dk.hex()}"


def verify_pin(pin, stored):
    try:
        _, iters, salt, digest = stored.split("$")
        dk = hashlib.pbkdf2_hmac("sha256", pin.encode(), bytes.fromhex(salt), int(iters))
        return hmac.compare_digest(dk.hex(), digest)
    except (ValueError, AttributeError):
        return False


def get_pin_hash(con):
    r = con.execute("SELECT value FROM settings WHERE key='_pin'").fetchone()
    return r["value"] if r else None


def set_pin(con, pin):
    con.execute("INSERT OR REPLACE INTO settings(key,value) VALUES('_pin', ?)", (hash_pin(pin),))


_login = {"fails": 0, "locked_until": 0.0}
_login_lock = threading.Lock()


def is_admin():
    ts = session.get("admin_ts")
    return bool(session.get("admin")) and ts is not None and time.time() - ts < ADMIN_IDLE_SECONDS


def admin_required(fn):
    @wraps(fn)
    def wrapper(*a, **kw):
        if not is_admin():
            session.pop("admin", None)
            return err("Admin login required.", 401, code="auth")
        session["admin_ts"] = time.time()  # idle timeout only counts admin actions
        return fn(*a, **kw)
    return wrapper


def err(message, status=400, **extra):
    return jsonify(error=message, **extra), status


def body():
    data = request.get_json(force=True, silent=True)
    return data if isinstance(data, dict) else {}


@app.before_request
def guard():
    if request.host not in ALLOWED_HOSTS:
        return err("Use http://127.0.0.1:8765 on this device.", 400)
    if request.method not in ("GET", "HEAD", "OPTIONS"):
        origin = request.headers.get("Origin")
        if origin and origin not in ALLOWED_ORIGINS:
            return err("Bad origin.", 403, code="origin")
        tok, sent = session.get("csrf"), request.headers.get("X-CSRF-Token", "")
        if not tok or not hmac.compare_digest(tok, sent):
            return err("Session expired. Please retry.", 403, code="csrf")


@app.after_request
def headers(resp):
    resp.headers["X-Content-Type-Options"] = "nosniff"
    resp.headers["X-Frame-Options"] = "DENY"
    resp.headers["Referrer-Policy"] = "no-referrer"
    resp.headers["Content-Security-Policy"] = (
        "default-src 'self'; script-src 'self' 'unsafe-eval'; style-src 'self'; "
        "img-src 'self' data:; media-src 'self' blob:; connect-src 'self'; "
        "object-src 'none'; base-uri 'none'; form-action 'none'; frame-ancestors 'none'")
    if request.path.startswith(("/api/", "/export/")):
        resp.headers["Cache-Control"] = "no-store"
    return resp


@app.errorhandler(Exception)
def on_error(e):
    if isinstance(e, HTTPException):
        if request.path.startswith("/api/"):
            return err(e.description, e.code)
        return e
    log.exception("Unhandled error")
    return err("Server error. See the Termux console.", 500)


@app.get("/")
def index():
    return render_template("index.html")


def kiosk_cfg(cfg):
    return dict(require_liveness=cfg["require_liveness"], confirm_frames=cfg["confirm_frames"],
                detect_size=cfg["detect_size"], min_face_pct=cfg["min_face_pct"])


@app.get("/api/session")
def get_session():
    session.permanent = True
    if "csrf" not in session:
        session["csrf"] = secrets.token_urlsafe(32)
    with db() as con:
        pin_set = get_pin_hash(con) is not None
        cfg = get_settings(con)
    return jsonify(csrf=session["csrf"], admin=is_admin(), pin_set=pin_set,
                   server_time=now().isoformat(timespec="seconds"),
                   kiosk=kiosk_cfg(cfg))


def _check_new_pin(pin):
    if not isinstance(pin, str) or len(pin) < MIN_PIN_LEN:
        raise ValueError(f"PIN must be at least {MIN_PIN_LEN} characters.")
    if len(set(pin)) == 1:
        raise ValueError("PIN is too simple.")


@app.post("/api/setup-pin")
def setup_pin():
    data = body()
    with db() as con:
        if get_pin_hash(con):
            return err("A PIN is already set.", 409)
        try:
            _check_new_pin(data.get("pin"))
        except ValueError as e:
            return err(str(e))
        set_pin(con, data["pin"])
        audit(con, "pin_setup")
    session["admin"], session["admin_ts"] = True, time.time()
    return jsonify(ok=True)


@app.post("/api/login")
def login():
    with _login_lock:
        wait = _login["locked_until"] - time.time()
        if wait > 0:
            return err(f"Too many attempts. Try again in {int(wait) + 1}s.", 429)
    pin = body().get("pin")
    with db() as con:
        stored = get_pin_hash(con)
        ok = bool(stored) and isinstance(pin, str) and verify_pin(pin, stored)
        with _login_lock:
            if ok:
                _login.update(fails=0, locked_until=0.0)
            else:
                _login["fails"] += 1
                if _login["fails"] >= 5:
                    _login["locked_until"] = time.time() + min(300, 15 * 2 ** (_login["fails"] - 5))
        audit(con, "login_ok" if ok else "login_failed")
    if not ok:
        return err("Wrong PIN.", 401)
    session["admin"], session["admin_ts"] = True, time.time()
    return jsonify(ok=True)


@app.post("/api/logout")
def logout():
    session.pop("admin", None)
    session.pop("admin_ts", None)
    return jsonify(ok=True)


@app.post("/api/change-pin")
@admin_required
def change_pin():
    data = body()
    with db() as con:
        if not verify_pin(str(data.get("old", "")), get_pin_hash(con) or ""):
            return err("Current PIN is wrong.", 403)
        try:
            _check_new_pin(data.get("new"))
        except ValueError as e:
            return err(str(e))
        set_pin(con, data["new"])
        audit(con, "pin_changed")
    return jsonify(ok=True)


# ------------------------------------------------------------ descriptors / roster
def validate_descriptor(d):
    if not isinstance(d, list) or len(d) != DESCRIPTOR_LEN:
        raise ValueError("Expected a 128-value face descriptor.")
    try:
        vals = [float(x) for x in d]
    except (TypeError, ValueError):
        raise ValueError("Invalid face descriptor.")
    if not all(math.isfinite(x) and abs(x) < 10 for x in vals):
        raise ValueError("Invalid face descriptor.")
    return vals


def validate_samples(raw, lo=3, hi=10):
    if not isinstance(raw, list) or not (lo <= len(raw) <= hi):
        raise ValueError(f"Provide {lo}-{hi} face samples.")
    samples = [validate_descriptor(s) for s in raw]
    worst = max(math.dist(a, b) for i, a in enumerate(samples) for b in samples[i + 1:])
    if worst > 0.50:  # near (~50 cm) and far (~1 m) samples of one person differ a little more
        raise ValueError("Samples do not look like one consistent face. Re-capture with one person in frame.")
    return samples


def student_distance(samples, d):
    """Mean of the two closest samples: steadier than min(), less strict than mean-of-all."""
    ds = sorted(math.dist(s, d) for s in samples)[:2]
    return sum(ds) / len(ds)


_cache = {"templates": None}
_cache_lock = threading.Lock()


def invalidate_roster():
    with _cache_lock:
        _cache["templates"] = None


def templates():
    with _cache_lock:
        if _cache["templates"] is None:
            with db() as con:
                rows = con.execute("SELECT id, student_no, name, descriptor FROM students WHERE active=1").fetchall()
            _cache["templates"] = [dict(id=r["id"], student_no=r["student_no"], name=r["name"],
                                        samples=parse_samples(r["descriptor"])) for r in rows]
        return _cache["templates"]


def find_duplicate_face(con, samples, threshold, exclude_id=None):
    for r in con.execute("SELECT id, student_no, name, descriptor FROM students").fetchall():
        if r["id"] == exclude_id:
            continue
        other = parse_samples(r["descriptor"])
        if min(student_distance(other, s) for s in samples) <= threshold:
            return r
    return None


def clean_text(value, field, maxlen, required=True):
    v = str(value or "").strip()
    if required and not v:
        raise ValueError(f"{field} is required.")
    if len(v) > maxlen:
        raise ValueError(f"{field} is too long (max {maxlen}).")
    return v


# ----------------------------------------------------------------- recognition
_streak = {"sid": None, "count": 0, "t": 0.0}
_streak_lock = threading.Lock()
STREAK_WINDOW = 4.0


def reset_streak():
    with _streak_lock:
        _streak.update(sid=None, count=0)


def auto_status(n, s):
    h, m = map(int, s["class_start"].split(":"))
    limit = datetime.combine(n.date(), datetime.min.time()).replace(hour=h, minute=m) + timedelta(minutes=s["late_after_min"])
    return "late" if n > limit else "present"


def match_verdict(scored, s):
    """scored = [(distance, template), ...] sorted ascending. -> 'unknown' | 'ambiguous' | 'match'."""
    best_d = scored[0][0]
    if best_d > s["threshold"]:
        return "unknown"
    if len(scored) > 1:
        second_d = scored[1][0]
        if second_d <= s["threshold"] and second_d - best_d < s["margin"]:
            return "ambiguous"
    return "match"


@app.get("/api/kiosk-config")
def kiosk_config():
    with db() as con:
        s = get_settings(con)
    return jsonify(kiosk_cfg(s))


@app.post("/api/recognize")
def recognize():
    """Match one face descriptor. Records attendance after N consistent frames."""
    data = body()
    try:
        d = validate_descriptor(data.get("descriptor"))
    except ValueError as e:
        return err(str(e))
    with db() as con:
        s = get_settings(con)
    tpls = templates()
    if not tpls:
        return jsonify(status="no_students")
    scored = sorted(((student_distance(t["samples"], d), t) for t in tpls), key=lambda x: x[0])
    best_d, best = scored[0]
    extra = {"distance": round(best_d, 3)} if is_admin() else {}
    verdict = match_verdict(scored, s)
    if verdict != "match":
        reset_streak()
        return jsonify(status=verdict, **extra)
    if s["require_liveness"] and not data.get("live"):
        return jsonify(status="need_liveness", **extra)
    t = time.monotonic()
    with _streak_lock:
        if _streak["sid"] == best["id"] and t - _streak["t"] <= STREAK_WINDOW:
            _streak["count"] += 1
        else:
            _streak.update(sid=best["id"], count=1)
        _streak["t"] = t
        count = _streak["count"]
        if count < s["confirm_frames"]:
            return jsonify(status="confirming", progress=count, needed=s["confirm_frames"], **extra)
        _streak.update(sid=None, count=0)
    n = now()
    day = n.date().isoformat()
    status = auto_status(n, s)
    with db() as con:
        cur = con.execute(
            "INSERT OR IGNORE INTO attendance(student_id, attend_date, time_in, method, status) VALUES(?,?,?,?,?)",
            (best["id"], day, n.strftime("%H:%M:%S"), "face", status))
        if cur.rowcount == 0:
            row = con.execute("SELECT time_in, status FROM attendance WHERE student_id=? AND attend_date=?",
                              (best["id"], day)).fetchone()
            return jsonify(status="duplicate", student_id=best["id"], name=best["name"],
                           student_no=best["student_no"], time_in=row["time_in"], att_status=row["status"], **extra)
    return jsonify(status="recorded", student_id=best["id"], name=best["name"], student_no=best["student_no"],
                   time_in=n.strftime("%H:%M:%S"), att_status=status, **extra)


@app.post("/api/calibration/match")
@admin_required
def calibration_match():
    """Admin-only probe used by the Calibration tab: reports distances but NEVER records attendance."""
    data = body()
    try:
        d = validate_descriptor(data.get("descriptor"))
    except ValueError as e:
        return err(str(e))
    with db() as con:
        s = get_settings(con)
    out = dict(threshold=s["threshold"], margin=s["margin"])
    tpls = templates()
    if not tpls:
        return jsonify(status="no_students", **out)
    scored = sorted(((student_distance(t["samples"], d), t) for t in tpls), key=lambda x: x[0])
    best_d, best = scored[0]
    out.update(status=match_verdict(scored, s), distance=round(best_d, 3),
               best=dict(id=best["id"], name=best["name"], student_no=best["student_no"]))
    if len(scored) > 1:
        out.update(second_distance=round(scored[1][0], 3), second_name=scored[1][1]["name"])
    try:
        target = int(data.get("target_id"))
    except (TypeError, ValueError):
        target = None
    if target is not None:
        for dist_, t in scored:
            if t["id"] == target:
                out.update(target_distance=round(dist_, 3), target_match=dist_ <= s["threshold"])
                break
    return jsonify(**out)


# ---------------------------------------------------------------------- students
@app.get("/api/students")
@admin_required
def list_students():
    with db() as con:
        rows = con.execute("SELECT id, student_no, name, section, active, descriptor FROM students ORDER BY name").fetchall()
    return jsonify([dict(id=r["id"], student_no=r["student_no"], name=r["name"], section=r["section"],
                         active=r["active"], samples=len(parse_samples(r["descriptor"]))) for r in rows])


@app.post("/api/students")
@admin_required
def add_student():
    data = body()
    try:
        student_no = clean_text(data.get("student_no"), "Student number", 32)
        name = clean_text(data.get("name"), "Name", 100)
        section = clean_text(data.get("section"), "Section", 50, required=False)
        samples = validate_samples(data.get("samples"))
    except ValueError as e:
        return err(str(e))
    try:
        with db() as con:
            s = get_settings(con)
            dup = find_duplicate_face(con, samples, s["threshold"])
            if dup:
                return err(f"This face already matches {dup['name']} ({dup['student_no']}). Not enrolled.", 409, code="dup_face")
            cur = con.execute(
                "INSERT INTO students(student_no,name,section,descriptor,created_at,updated_at) VALUES(?,?,?,?,?,?)",
                (student_no, name, section, json.dumps(samples), now().isoformat(timespec="seconds"),
                 now().isoformat(timespec="seconds")))
            audit(con, "student_added", f"{student_no} {name}")
    except sqlite3.IntegrityError:
        return err("That student number already exists.", 409)
    invalidate_roster()
    return jsonify(ok=True, id=cur.lastrowid)


@app.put("/api/students/<int:sid>")
@admin_required
def edit_student(sid):
    data = body()
    try:
        student_no = clean_text(data.get("student_no"), "Student number", 32)
        name = clean_text(data.get("name"), "Name", 100)
        section = clean_text(data.get("section"), "Section", 50, required=False)
    except ValueError as e:
        return err(str(e))
    try:
        with db() as con:
            cur = con.execute("UPDATE students SET student_no=?, name=?, section=?, updated_at=? WHERE id=?",
                              (student_no, name, section, now().isoformat(timespec="seconds"), sid))
            if cur.rowcount == 0:
                return err("Student not found.", 404)
            audit(con, "student_edited", f"id={sid} {student_no} {name}")
    except sqlite3.IntegrityError:
        return err("That student number already exists.", 409)
    invalidate_roster()
    return jsonify(ok=True)


@app.post("/api/students/<int:sid>/samples")
@admin_required
def update_samples(sid):
    """Re-enrol (mode=replace) or add extra samples (mode=add, keeps newest 10)."""
    data = body()
    try:
        new = validate_samples(data.get("samples"), lo=3 if data.get("mode") != "add" else 1)
    except ValueError as e:
        return err(str(e))
    with db() as con:
        row = con.execute("SELECT descriptor FROM students WHERE id=?", (sid,)).fetchone()
        if not row:
            return err("Student not found.", 404)
        dup = find_duplicate_face(con, new, get_settings(con)["threshold"], exclude_id=sid)
        if dup:
            return err(f"This face already matches {dup['name']} ({dup['student_no']}).", 409, code="dup_face")
        samples = (parse_samples(row["descriptor"]) + new)[-10:] if data.get("mode") == "add" else new
        con.execute("UPDATE students SET descriptor=?, updated_at=? WHERE id=?",
                    (json.dumps(samples), now().isoformat(timespec="seconds"), sid))
        audit(con, "samples_updated", f"id={sid} mode={data.get('mode', 'replace')}")
    invalidate_roster()
    return jsonify(ok=True)


@app.post("/api/students/<int:sid>/toggle")
@admin_required
def toggle_student(sid):
    with db() as con:
        row = con.execute("SELECT active, student_no FROM students WHERE id=?", (sid,)).fetchone()
        if not row:
            return err("Student not found.", 404)
        con.execute("UPDATE students SET active=? WHERE id=?", (0 if row["active"] else 1, sid))
        audit(con, "student_toggled", f"{row['student_no']} active={0 if row['active'] else 1}")
    invalidate_roster()
    return jsonify(ok=True)


@app.delete("/api/students/<int:sid>")
@admin_required
def delete_student(sid):
    with db() as con:
        row = con.execute("SELECT student_no, name FROM students WHERE id=?", (sid,)).fetchone()
        if not row:
            return err("Student not found.", 404)
        con.execute("DELETE FROM attendance WHERE student_id=?", (sid,))
        con.execute("DELETE FROM students WHERE id=?", (sid,))
        audit(con, "student_deleted", f"{row['student_no']} {row['name']}")
    invalidate_roster()
    return jsonify(ok=True)


# -------------------------------------------------------------------- attendance
def parse_day(value, default=None):
    if not value:
        return default
    return date.fromisoformat(str(value))


@app.post("/api/attendance/manual")
@admin_required
def manual_mark():
    data = body()
    try:
        sid = int(data.get("student_id"))
        status = data.get("status", "present")
        if status not in ("present", "late", "excused"):
            raise ValueError("Status must be present, late or excused.")
        day = parse_day(data.get("date"), now().date())
        if day > now().date():
            raise ValueError("Cannot record attendance in the future.")
        note = clean_text(data.get("note"), "Note", 200, required=False)
    except (TypeError, ValueError) as e:
        return err(str(e) if str(e) else "Invalid request.")
    n = now()
    time_in = n.strftime("%H:%M:%S") if status != "excused" and day == n.date() else ""
    with db() as con:
        st = con.execute("SELECT student_no, name FROM students WHERE id=?", (sid,)).fetchone()
        if not st:
            return err("Student not found.", 404)
        con.execute("""INSERT INTO attendance(student_id, attend_date, time_in, method, status, note)
                       VALUES(?,?,?,'manual',?,?)
                       ON CONFLICT(student_id, attend_date) DO UPDATE SET
                         time_in=excluded.time_in, method='manual', status=excluded.status, note=excluded.note""",
                    (sid, day.isoformat(), time_in, status, note))
        audit(con, "manual_mark", f"{st['student_no']} {day} {status} {note}")
    return jsonify(ok=True, time_in=time_in)


@app.delete("/api/attendance/<int:aid>")
@admin_required
def delete_attendance(aid):
    with db() as con:
        row = con.execute("""SELECT a.attend_date, a.status, s.student_no FROM attendance a
                             JOIN students s ON s.id=a.student_id WHERE a.id=?""", (aid,)).fetchone()
        if not row:
            return err("Record not found.", 404)
        con.execute("DELETE FROM attendance WHERE id=?", (aid,))
        audit(con, "attendance_removed", f"{row['student_no']} {row['attend_date']} {row['status']}")
    return jsonify(ok=True)


# ----------------------------------------------------------------------- reports
def holidays_between(con, start, end):
    return {r["day"] for r in con.execute("SELECT day FROM holidays WHERE day>=? AND day<=?",
                                          (start.isoformat(), end.isoformat()))}


def is_school_day(con, d, s):
    return d.isoweekday() in s["school_days"] and not holidays_between(con, d, d)


def daily_rows(con, day):
    s = get_settings(con)
    school = is_school_day(con, day, s)
    rows = con.execute("""
        SELECT s.id, s.student_no, s.name, s.section, s.active,
               a.id AS att_id, a.time_in, a.method, a.status AS att_status, a.note
        FROM students s LEFT JOIN attendance a ON a.student_id=s.id AND a.attend_date=?
        ORDER BY s.name""", (day.isoformat(),)).fetchall()
    out = []
    for r in rows:
        d = dict(r)
        if d["att_status"]:
            d["status"] = d["att_status"].capitalize()
        elif not d["active"]:
            d["status"] = "Inactive"
        elif day > now().date():
            d["status"] = "Upcoming"
        elif not school:
            d["status"] = "No class"
        else:
            d["status"] = "Absent"
        out.append(d)
    return out, school


def daily_totals(rows):
    keys = ("Present", "Late", "Excused", "Absent", "Inactive")
    return {k.lower(): sum(1 for r in rows if r["status"] == k) for k in keys}


def month_bounds(month):
    try:
        y, m = map(int, month.split("-"))
        first = date(y, m, 1)
    except (ValueError, AttributeError):
        raise ValueError("Month must look like 2026-10.")
    nxt = date(y + (m == 12), m % 12 + 1, 1)
    return first, nxt - timedelta(days=1)


def _month_data(con, month):
    """Shared by the on-screen table, the CSV and the Excel export so all three always agree."""
    s = get_settings(con)
    first, last = month_bounds(month)
    end = min(last, now().date())
    hol = holidays_between(con, first, last)
    days, d = [], first
    while d <= end:
        if d.isoweekday() in s["school_days"] and d.isoformat() not in hol:
            days.append(d.isoformat())
        d += timedelta(days=1)
    recs = {}
    for r in con.execute("SELECT student_id, attend_date, status FROM attendance WHERE attend_date>=? AND attend_date<=?",
                         (first.isoformat(), last.isoformat())):
        recs.setdefault(r["student_id"], {})[r["attend_date"]] = r["status"]
    students = con.execute("SELECT id, student_no, name, section, active FROM students ORDER BY name").fetchall()
    return days, recs, students


def monthly_rows(con, month):
    days, recs, students = _month_data(con, month)
    school_days = set(days)
    out = []
    for st in students:
        mine = {k: v for k, v in recs.get(st["id"], {}).items() if k in school_days}
        if not st["active"] and not recs.get(st["id"]):
            continue
        present = sum(1 for v in mine.values() if v == "present")
        late = sum(1 for v in mine.values() if v == "late")
        excused = sum(1 for v in mine.values() if v == "excused")
        absent = len(school_days) - present - late - excused
        denom = len(school_days) - excused
        out.append(dict(student_no=st["student_no"], name=st["name"], section=st["section"], active=st["active"],
                        present=present, late=late, excused=excused, absent=absent,
                        rate=round(100.0 * (present + late) / denom, 1) if denom > 0 else None))
    return out, len(school_days)


def monthly_grid(con, month):
    """Day-by-day status per student: ([iso days], [{student_no, name, section, marks: {day: status}}])."""
    days, recs, students = _month_data(con, month)
    out = []
    for st in students:
        if not st["active"] and not recs.get(st["id"]):
            continue
        mine = recs.get(st["id"], {})
        out.append(dict(student_no=st["student_no"], name=st["name"], section=st["section"],
                        marks={d: mine.get(d, "absent") for d in days}))
    return days, out


@app.get("/api/report")
@admin_required
def report():
    try:
        day = parse_day(request.args.get("date"), now().date())
    except ValueError:
        return err("Invalid date.")
    with db() as con:
        rows, school = daily_rows(con, day)
    return jsonify(date=day.isoformat(), school_day=school, rows=rows, totals=daily_totals(rows))


@app.get("/api/monthly")
@admin_required
def monthly():
    month = request.args.get("month", now().strftime("%Y-%m"))
    try:
        with db() as con:
            rows, n_days = monthly_rows(con, month)
    except ValueError as e:
        return err(str(e))
    return jsonify(month=month, school_days=n_days, rows=rows)


def csv_safe(v):
    """Neutralise spreadsheet formula injection in user-controlled text."""
    s = "" if v is None else str(v)
    return "'" + s if s[:1] in ("=", "+", "-", "@", "\t", "\r") else s


def csv_response(rows, filename):
    out = io.StringIO()
    csv.writer(out).writerows(rows)
    return Response("\ufeff" + out.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": f"attachment; filename={filename}"})


@app.get("/export/daily.csv")
@admin_required
def export_daily():
    try:
        day = parse_day(request.args.get("date"), now().date())
    except ValueError:
        return err("Invalid date.")
    with db() as con:
        rows, school = daily_rows(con, day)
    data = [["Date", day.isoformat()], ["Student Number", "Student Name", "Section", "Status", "Time In", "Method", "Note"]]
    data += [[csv_safe(r["student_no"]), csv_safe(r["name"]), csv_safe(r["section"]), r["status"],
              r["time_in"] or "", r["method"] or "", csv_safe(r["note"])] for r in rows]
    return csv_response(data, f"attendance_{day.isoformat()}.csv")


@app.get("/export/monthly.csv")
@admin_required
def export_monthly():
    month = request.args.get("month", now().strftime("%Y-%m"))
    try:
        with db() as con:
            rows, n_days = monthly_rows(con, month)
    except ValueError as e:
        return err(str(e))
    data = [["Month", month, "School days so far", n_days],
            ["Student Number", "Student Name", "Section", "Present", "Late", "Excused", "Absent", "Attendance %"]]
    data += [[csv_safe(r["student_no"]), csv_safe(r["name"]), csv_safe(r["section"]), r["present"], r["late"],
              r["excused"], r["absent"], "" if r["rate"] is None else r["rate"]] for r in rows]
    return csv_response(data, f"attendance_month_{month}.csv")


# ------------------------------------------------------------------ Excel (.xlsx)
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
STATUS_FILL = {"present": "E3F4EA", "late": "FFF0CC", "excused": "E8E8FB", "absent": "FDE8E8"}
STATUS_CODE = {"present": "P", "late": "L", "excused": "E", "absent": "A"}


def _xl():
    """openpyxl is imported lazily so the kiosk still runs if it is not installed."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None
    return openpyxl, Alignment, Border, Font, PatternFill, Side, get_column_letter


def _xl_missing():
    return err("Excel export needs the openpyxl package. In Termux run: python -m pip install openpyxl", 501)


class _Sheet:
    """Small helper: text is always stored as text, never as a formula (spreadsheet-injection safe)."""

    def __init__(self, ws, xl):
        self.ws = ws
        _, self.Alignment, self.Border, self.Font, self.PatternFill, self.Side, self.col = xl
        side = self.Side(style="thin", color="CBD2DC")
        self.border = self.Border(left=side, right=side, top=side, bottom=side)

    def put(self, row, col, value, bold=False, fill=None, color=None, fmt=None, align=None, wrap=False, border=False, size=None):
        c = self.ws.cell(row=row, column=col, value=value)
        if isinstance(value, str):
            c.data_type = "s"  # even if it starts with = + - @
        if bold or color or size:
            c.font = self.Font(bold=bold, color=color, size=size)
        if fill:
            c.fill = self.PatternFill("solid", start_color=fill, end_color=fill)
        if fmt:
            c.number_format = fmt
        if align or wrap:
            c.alignment = self.Alignment(horizontal=align, vertical="center", wrap_text=wrap)
        if border:
            c.border = self.border
        return c

    def info(self, row, label, value):
        """Label merged across A:B (so long labels are not cut off), value in C."""
        self.ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        self.put(row, 1, label, bold=True)
        self.put(row, 3, value, align="left")

    def header(self, row, labels, start_col=1):
        for i, text in enumerate(labels):
            self.put(row, start_col + i, text, bold=True, fill="174EA6", color="FFFFFF", align="center", wrap=True, border=True)

    def widths(self, widths):
        for i, w in enumerate(widths, 1):
            self.ws.column_dimensions[self.col(i)].width = w

    def print_setup(self):
        self.ws.page_setup.orientation = "landscape"
        self.ws.page_setup.fitToWidth = 1
        self.ws.page_setup.fitToHeight = 0
        self.ws.sheet_properties.pageSetUpPr.fitToPage = True


def _rate_fill(rate):
    if rate is None:
        return None
    return "E3F4EA" if rate >= 90 else "FFF0CC" if rate >= 75 else "FDE8E8"


def xlsx_response(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), mimetype=XLSX_MIME,
                    headers={"Content-Disposition": f"attachment; filename={filename}"})


@app.get("/export/daily.xlsx")
@admin_required
def export_daily_xlsx():
    xl = _xl()
    if not xl:
        return _xl_missing()
    try:
        day = parse_day(request.args.get("date"), now().date())
    except ValueError:
        return err("Invalid date.")
    with db() as con:
        rows, school = daily_rows(con, day)
    tot = daily_totals(rows)
    wb = xl[0].Workbook()
    sh = _Sheet(wb.active, xl)
    sh.ws.title = "Daily Attendance"
    sh.put(1, 1, "Daily Attendance Report", bold=True, size=14)
    sh.info(2, "Date", day.strftime("%A, %B %d, %Y"))
    sh.info(3, "School day", "Yes" if school else "No (weekend / holiday)")
    sh.info(4, "Totals", f"{tot['present']} present • {tot['late']} late • {tot['excused']} excused • "
                         f"{tot['absent']} absent • {tot['inactive']} inactive")
    sh.info(5, "Generated", now().strftime("%Y-%m-%d %H:%M"))
    sh.header(7, ["No.", "Student Number", "Student Name", "Section", "Status", "Time In", "Method", "Note"])
    r = 8
    for i, row in enumerate(rows, 1):
        fill = STATUS_FILL.get(row["status"].lower())
        sh.put(r, 1, i, border=True, align="center")
        sh.put(r, 2, row["student_no"], border=True)
        sh.put(r, 3, row["name"], border=True)
        sh.put(r, 4, row["section"] or "", border=True)
        sh.put(r, 5, row["status"], border=True, fill=fill, bold=True, align="center")
        sh.put(r, 6, row["time_in"] or "", border=True, align="center")
        sh.put(r, 7, row["method"] or "", border=True, align="center")
        sh.put(r, 8, row["note"] or "", border=True)
        r += 1
    sh.widths([6, 16, 30, 16, 12, 11, 11, 34])
    sh.ws.freeze_panes = "A8"
    sh.print_setup()
    return xlsx_response(wb, f"attendance_{day.isoformat()}.xlsx")


@app.get("/export/monthly.xlsx")
@admin_required
def export_monthly_xlsx():
    xl = _xl()
    if not xl:
        return _xl_missing()
    month = request.args.get("month", now().strftime("%Y-%m"))
    try:
        with db() as con:
            rows, n_days = monthly_rows(con, month)
            days, grid = monthly_grid(con, month)
        first, _ = month_bounds(month)
    except ValueError as e:
        return err(str(e))
    wb = xl[0].Workbook()

    # ---- Sheet 1: monthly summary ------------------------------------------------
    sh = _Sheet(wb.active, xl)
    sh.ws.title = "Monthly Summary"
    sh.put(1, 1, "Monthly Attendance Summary", bold=True, size=14)
    sh.info(2, "Month", first.strftime("%B %Y"))
    sh.info(3, "School days counted", n_days)
    sh.put(3, 4, "(weekends, holidays and future days excluded)")
    sh.info(4, "Generated", now().strftime("%Y-%m-%d %H:%M"))
    sh.info(5, "Rule", "Attendance % = (Present + Late) ÷ (School days − Excused)")
    sh.header(7, ["No.", "Student Number", "Student Name", "Section", "Present", "Late", "Excused", "Absent", "Attendance %"])
    r = 8
    for i, row in enumerate(rows, 1):
        sh.put(r, 1, i, border=True, align="center")
        sh.put(r, 2, row["student_no"], border=True)
        sh.put(r, 3, row["name"], border=True)
        sh.put(r, 4, row["section"] or "", border=True)
        sh.put(r, 5, row["present"], border=True, align="center", fill=STATUS_FILL["present"] if row["present"] else None)
        sh.put(r, 6, row["late"], border=True, align="center", fill=STATUS_FILL["late"] if row["late"] else None)
        sh.put(r, 7, row["excused"], border=True, align="center", fill=STATUS_FILL["excused"] if row["excused"] else None)
        sh.put(r, 8, row["absent"], border=True, align="center", fill=STATUS_FILL["absent"] if row["absent"] else None)
        rate = row["rate"]
        sh.put(r, 9, "" if rate is None else rate / 100.0, border=True, align="center", bold=True,
               fmt="0.0%", fill=_rate_fill(rate))
        r += 1
    if rows:
        tp, tl, te, ta = (sum(x[k] for x in rows) for k in ("present", "late", "excused", "absent"))
        denom = sum(n_days - x["excused"] for x in rows)
        class_rate = (tp + tl) / denom if denom > 0 else None
        r += 1
        sh.put(r, 3, "Class total", bold=True, border=True, fill="E4E9F0")
        sh.put(r, 4, "", border=True, fill="E4E9F0")
        for col, v in zip(range(5, 9), (tp, tl, te, ta)):
            sh.put(r, col, v, bold=True, border=True, align="center", fill="E4E9F0")
        sh.put(r, 9, "" if class_rate is None else class_rate, bold=True, border=True, align="center",
               fmt="0.0%", fill="E4E9F0")
    sh.widths([6, 16, 30, 16, 10, 10, 10, 10, 14])
    sh.ws.freeze_panes = "A8"
    sh.print_setup()

    # ---- Sheet 2: day by day grid ------------------------------------------------
    g = _Sheet(wb.create_sheet("Day by Day"), xl)
    g.put(1, 1, "Day-by-Day Attendance", bold=True, size=14)
    g.put(2, 1, "Legend", bold=True)
    g.put(2, 2, "P = Present     L = Late     E = Excused     A = Absent")
    labels = [date.fromisoformat(d) for d in days]
    g.header(4, ["Student Number", "Student Name", "Section"] + [f"{d:%a}\n{d.day}" for d in labels] + ["P", "L", "E", "A"])
    g.ws.row_dimensions[4].height = 32
    r = 5
    for row in grid:
        g.put(r, 1, row["student_no"], border=True)
        g.put(r, 2, row["name"], border=True)
        g.put(r, 3, row["section"] or "", border=True)
        counts = dict.fromkeys(STATUS_CODE, 0)
        for col, d in enumerate(days, 4):
            st = row["marks"][d]
            counts[st] += 1
            g.put(r, col, STATUS_CODE[st], align="center", border=True, fill=STATUS_FILL[st])
        for col, key in enumerate(("present", "late", "excused", "absent"), 4 + len(days)):
            g.put(r, col, counts[key], align="center", border=True, bold=True)
        r += 1
    g.widths([16, 30, 14] + [5.5] * len(days) + [5, 5, 5, 5])
    g.ws.freeze_panes = "D5"
    g.print_setup()

    # ---- Sheet 3: class totals per day -------------------------------------------
    t = _Sheet(wb.create_sheet("Totals per Day"), xl)
    t.put(1, 1, "Class Totals per School Day", bold=True, size=14)
    t.header(3, ["Date", "Day", "Present", "Late", "Excused", "Absent", "Attendance %"])
    r = 4
    for d, dd in zip(days, labels):
        c = dict.fromkeys(STATUS_CODE, 0)
        for row in grid:
            c[row["marks"][d]] += 1
        denom = len(grid) - c["excused"]
        t.put(r, 1, d, border=True, align="center")
        t.put(r, 2, f"{dd:%A}", border=True)
        for col, key in enumerate(("present", "late", "excused", "absent"), 3):
            t.put(r, col, c[key], border=True, align="center")
        rate = (c["present"] + c["late"]) / denom if denom > 0 else None
        t.put(r, 7, "" if rate is None else rate, border=True, align="center", bold=True, fmt="0.0%",
              fill=_rate_fill(None if rate is None else rate * 100))
        r += 1
    t.widths([13, 14, 10, 10, 10, 10, 14])
    t.ws.freeze_panes = "A4"
    t.print_setup()
    return xlsx_response(wb, f"attendance_summary_{month}.xlsx")


# --------------------------------------------------------- settings / calendar / ops
@app.get("/api/settings")
@admin_required
def read_settings():
    with db() as con:
        return jsonify(settings_for_api(get_settings(con)))


@app.put("/api/settings")
@admin_required
def write_settings():
    data, clean = body(), {}
    for k, v in data.items():
        if k not in SETTINGS:
            return err(f"Unknown setting: {k}")
        try:
            clean[k] = coerce_setting(k, v)
        except (TypeError, ValueError) as e:
            return err(str(e) or f"Invalid value for {k}.")
    with db() as con:
        for k, v in clean.items():
            con.execute("INSERT OR REPLACE INTO settings(key,value) VALUES(?,?)", (k, v))
        audit(con, "settings_changed", json.dumps(clean))
    return jsonify(ok=True)


@app.get("/api/holidays")
@admin_required
def list_holidays():
    with db() as con:
        rows = con.execute("SELECT day, reason FROM holidays ORDER BY day DESC LIMIT 400").fetchall()
    return jsonify([dict(r) for r in rows])


@app.post("/api/holidays")
@admin_required
def add_holiday():
    data = body()
    try:
        start = parse_day(data.get("start"))
        end = parse_day(data.get("end"), start)
        if not start or end < start or (end - start).days > 366:
            raise ValueError("Provide a valid start (and optional end) date, max 1 year.")
        reason = clean_text(data.get("reason"), "Reason", 100, required=False)
    except ValueError as e:
        return err(str(e))
    with db() as con:
        d = start
        while d <= end:
            con.execute("INSERT OR REPLACE INTO holidays(day, reason) VALUES(?,?)", (d.isoformat(), reason))
            d += timedelta(days=1)
        audit(con, "holiday_added", f"{start}..{end} {reason}")
    return jsonify(ok=True)


@app.delete("/api/holidays/<day>")
@admin_required
def delete_holiday(day):
    try:
        day = date.fromisoformat(day).isoformat()
    except ValueError:
        return err("Invalid date.")
    with db() as con:
        con.execute("DELETE FROM holidays WHERE day=?", (day,))
        audit(con, "holiday_removed", day)
    return jsonify(ok=True)


@app.get("/api/audit")
@admin_required
def audit_log():
    with db() as con:
        rows = con.execute("SELECT ts, action, detail FROM audit_log ORDER BY id DESC LIMIT 200").fetchall()
    return jsonify([dict(r) for r in rows])


@app.get("/api/backup")
@admin_required
def backup():
    """Consistent snapshot via SQLite's online backup API. Contains biometric templates: keep it encrypted."""
    fd, tmp = tempfile.mkstemp(suffix=".db")
    os.close(fd)
    try:
        src, dst = sqlite3.connect(DB_PATH), sqlite3.connect(tmp)
        src.backup(dst)
        dst.close()
        src.close()
        data = Path(tmp).read_bytes()
    finally:
        os.unlink(tmp)
    with db() as con:
        audit(con, "backup_downloaded")
    return Response(data, mimetype="application/octet-stream", headers={
        "Content-Disposition": f"attachment; filename=attendance_backup_{now().strftime('%Y%m%d_%H%M')}.db"})


@app.post("/api/purge")
@admin_required
def purge():
    """Retention: delete attendance records older than a date."""
    data = body()
    try:
        before = parse_day(data.get("before"))
        if not before or before > now().date() - timedelta(days=30):
            raise ValueError("Pick a date at least 30 days in the past.")
    except ValueError as e:
        return err(str(e))
    if data.get("confirm") != "PURGE":
        return err("Type PURGE to confirm.")
    with db() as con:
        cur = con.execute("DELETE FROM attendance WHERE attend_date<?", (before.isoformat(),))
        audit(con, "purge", f"before={before} deleted={cur.rowcount}")
    return jsonify(ok=True, deleted=cur.rowcount)


@app.get("/health")
def health():
    return jsonify(ok=True)


# --------------------------------------------------------------------------- main
def cli_set_pin():
    import getpass
    init_db()
    pin = getpass.getpass("New admin PIN: ")
    if pin != getpass.getpass("Repeat PIN: "):
        sys.exit("PINs do not match.")
    try:
        _check_new_pin(pin)
    except ValueError as e:
        sys.exit(str(e))
    with db() as con:
        set_pin(con, pin)
        audit(con, "pin_reset_cli")
    print("PIN updated.")


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    if "--set-pin" in sys.argv:
        cli_set_pin()
        sys.exit(0)
    init_db()
    print(f"Server time: {now():%Y-%m-%d %H:%M:%S}  (timezone: {time.tzname[0]}). "
          f"If this is wrong, run: export TZ=Asia/Manila")
    print(f"Open http://127.0.0.1:{PORT} in Chrome on this device.")
    # Bind only to this phone. Never change to 0.0.0.0 for a classroom deployment.
    app.run(host=HOST, port=PORT, debug=False, threaded=True)
